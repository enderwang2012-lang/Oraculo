#!/usr/bin/env python3
"""Promote the user's v4 selections into the corpus source and metadata."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from reset_corpus_review import (
    EDITORIAL_THEME_TO_EMOTION,
    cadence_group,
    editorial_theme,
    keep_reason,
    semantic_cluster,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "starbucks_now_passphrases.csv"
CANDIDATES = ROOT / "review" / "corpus_candidates_2026_10_12_full_v4.csv"
WORKBOOK = (
    ROOT
    / "outputs"
    / "01a11f97-376f-7fe2-8992-269c2e11fc0b"
    / "corpus_candidates_2026_10_12_full_v4_review.xlsx"
)
EN_MAP = ROOT / "scripts" / "phrases_en.json"
REVIEW_JSON = ROOT / "config" / "phrase_editorial_review.json"
REVIEW_CSV = ROOT / "review" / "corpus_review_2026_08_full.csv"
DISPATCH_OVERRIDES = ROOT / "config" / "phrase_dispatch_overrides.json"
SELECTED_SNAPSHOT = ROOT / "review" / "corpus_candidates_2026_10_12_full_v4_selected.csv"

SOURCE_FIELDS = ["id", "phrase", "approx_period", "theme", "source_id", "evidence", "notes"]
V4_FIELDS = [
    "candidate_id",
    "group",
    "phrase",
    "user_feedback",
    "review_decision",
    "status",
    "occasion",
    "window_hint",
    "source_kind",
    "source_work",
    "source_creator",
    "source_ref",
    "relation_kind",
    "review_note",
    "previous_id",
]

TRANSLATIONS = {
    "想入非非": "Let your thoughts run wild",
    "心血来潮": "On a whim",
    "心照不宣": "Understood without a word",
    "欲说还休": "Wanting to speak, then stopping",
    "自得其乐": "Finding joy in your own way",
    "随心所欲": "Do as you please",
    "天马行空": "Let your imagination take flight",
    "偷着乐": "Enjoy it on the sly",
    "有点意思": "Interesting",
    "刚刚好": "Just right",
    "我乐意": "Gladly",
    "好说好说": "Consider it handled",
    "来日方长": "There is time",
    "乐在其中": "Finding joy in it",
    "不亦乐乎": "What fun",
    "得闲": "At leisure",
    "自在": "At ease",
    "幸会": "A pleasure to meet you",
    "如愿": "As wished",
    "好耶": "Yay",
    "山河可爱": "The land is lovely",
    "大好河山": "A beautiful land",
    "国泰民安": "Peace for the country, peace for its people",
    "长长久久": "For a long, long time",
    "久久安康": "Long-lasting good health",
    "岁岁登高": "Rise higher every year",
    "不给糖就捣蛋": "Trick or treat",
    "装神弄鬼": "Playing dress-up",
    "谢谢你呀": "Thank you",
    "多谢关照": "Thanks for your kindness",
    "幸好有你": "Lucky to have you",
    "团团圆圆": "Together and whole",
    "冬至大如年": "The winter solstice is a festival of its own",
    "苹果分你一半": "Half an apple is yours",
    "平安无事": "Safe and sound",
    "叮叮当": "Jingle bells",
    "圣诞快乐": "Merry Christmas",
    "明年见": "See you next year",
    "来年可期": "The year ahead is full of promise",
    "旧岁再见": "Goodbye, old year",
    "暗香盈袖": "Fragrance fills the sleeves",
    "还来就菊花": "Come back for chrysanthemums",
    "红于二月花": "Redder than spring flowers",
    "橙黄橘绿时": "When oranges glow and tangerines turn green",
    "秋日胜春朝": "Autumn days beat spring mornings",
    "November Rain": "November Rain",
    "冬天快乐": "Happy winter",
    "能饮一杯无": "Care for a cup",
    "红泥小火炉": "A little stove of red clay",
    "not so bad": "Not so bad",
    "春又来": "Spring comes again",
    "All I want": "All I want",
    "Let it snow": "Let it snow",
    "Happy New Year": "Happy New Year",
    "明年会更好": "Next year will be better",
    "随遇而安": "Take things as they come",
    "云淡风轻": "Light clouds, gentle wind",
    "小确幸": "A little everyday joy",
    "自成一派": "A world of your own",
    "不妨一试": "Might as well try",
    "说来话长": "That is a long story",
    "好事发生": "Good things happen",
}

CADENCE_OVERRIDES = {
    "想入非非": "mood_state",
    "心血来潮": "mood_state",
    "欲说还休": "mood_state",
    "自得其乐": "medium_state",
    "随心所欲": "stance_short",
    "天马行空": "mood_state",
    "偷着乐": "mood_state",
    "有点意思": "mood_state",
    "刚刚好": "medium_state",
    "我乐意": "stance_short",
    "好说好说": "relational_phrase",
    "来日方长": "time_observation",
    "乐在其中": "medium_state",
    "得闲": "rest_state",
    "幸会": "relational_phrase",
    "如愿": "blessing_parallel",
    "好耶": "mood_state",
    "国泰民安": "blessing_parallel",
    "长长久久": "blessing_parallel",
    "久久安康": "blessing_parallel",
    "岁岁登高": "blessing_parallel",
    "装神弄鬼": "wordplay_social",
    "随遇而安": "rest_state",
    "小确幸": "mood_state",
    "自成一派": "medium_state",
    "说来话长": "relational_phrase",
}


def load_selected() -> list[dict[str, str]]:
    with CANDIDATES.open(encoding="utf-8", newline="") as handle:
        candidate_rows = list(csv.DictReader(handle))
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheet = workbook["候选评审"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(values_only=True)]
    decisions = {
        str(row["candidate_id"]): str(row["review_decision"] or "")
        for row in rows
        if row.get("candidate_id")
    }
    selected = [
        row for row in candidate_rows if decisions.get(row["candidate_id"]) == "通过"
    ]
    if len(selected) != 62:
        raise SystemExit(f"Expected 62 selected candidates, found {len(selected)}")
    if len(set(row["phrase"] for row in selected)) != len(selected):
        raise SystemExit("Selected candidates contain duplicate phrases")
    return selected


def source_metadata(row: dict[str, str]) -> tuple[str, str, str, str]:
    group = row["group"]
    phrase = row["phrase"]
    source_kind = row["source_kind"]
    source_ref = row["source_ref"]
    note = f"2026-10-12 v4；用户评审通过。{row['review_note']}"

    if group == "general":
        return "生活态度", "GEN9", "generated", note
    if group == "festival":
        return "节日祝福", "GEN9", "generated", note
    if source_kind == "classical_excerpt":
        return "诗性意象", source_ref, "official_text", note
    if phrase == "明年会更好":
        return "新年祝福", source_ref, "generated", note
    return "歌曲与季节", source_ref, "official_text", note


def dispatch_override_for(row: dict[str, str], pid: str) -> dict | None:
    phrase = row["phrase"]
    group = row["group"]
    only_when: list[str] = []

    if phrase in {"山河可爱", "大好河山", "国泰民安"}:
        only_when = ["festival:national_day"]
    elif phrase in {"团团圆圆", "冬至大如年", "春又来"}:
        only_when = ["solar_term:dongzhi"]
    elif phrase in {"苹果分你一半", "平安无事", "叮叮当", "圣诞快乐", "All I want", "Let it snow"}:
        only_when = ["festival:christmas"]
    elif phrase in {"明年见", "来年可期", "旧岁再见", "Happy New Year", "明年会更好"}:
        only_when = ["festival:new_year"]
    elif phrase == "November Rain":
        only_when = ["month:11"]
    elif phrase == "冬天快乐":
        only_when = ["season:winter"]
    elif phrase in {"红于二月花", "橙黄橘绿时", "秋日胜春朝"}:
        only_when = ["season:autumn"]
    elif phrase in {"能饮一杯无", "红泥小火炉"}:
        only_when = ["solar_term:xiaoxue"]

    if not only_when:
        return None
    return {"universal": False, "onlyWhen": only_when, "boost": []}


def editorial_record(
    row: dict[str, str],
    pid: str,
    source_theme: str,
    source_id: str,
    evidence: str,
) -> dict:
    phrase = row["phrase"]
    theme = editorial_theme(phrase, source_theme)
    rights_status = (
        "editorial_origin_recorded"
        if evidence == "generated"
        else "external_source_recorded_no_clearance_claim"
    )
    return {
        "phrase": phrase,
        "sourceTheme": source_theme,
        "evidence": evidence,
        "sourceId": source_id,
        "sourceCategory": "editorial_original" if evidence == "generated" else "external_observed",
        "rightsStatus": rights_status,
        "reviewBasis": "fresh_initial_content_review",
        "decision": "keep",
        "reviewCode": "approved_initial",
        "lifecycle": "active",
        "proposedPhrase": "",
        "proposedEnglish": "",
        "editorialTheme": theme,
        "emotionTheme": EDITORIAL_THEME_TO_EMOTION[theme],
        "semanticCluster": semantic_cluster(phrase, source_theme),
        "cadenceGroup": CADENCE_OVERRIDES.get(
            phrase,
            cadence_group(phrase, source_theme),
        ),
        "english": TRANSLATIONS[phrase],
        "englishStatus": "approved",
        "dispatchStatus": "manual" if dispatch_override_for(row, pid) else "rule",
        "dispatchConstraint": dispatch_override_for(row, pid)["onlyWhen"] if dispatch_override_for(row, pid) else [],
        "reviewerNote": row["review_note"],
        "reviewReason": keep_reason(pid, theme),
    }


def main() -> None:
    selected = load_selected()

    with SOURCE.open(encoding="utf-8-sig", newline="") as handle:
        source_rows = list(csv.DictReader(handle))
    existing_phrases = {row["phrase"].strip() for row in source_rows}
    existing_ids = {int(row["id"]) for row in source_rows}
    next_id = max(existing_ids) + 1
    for row in selected:
        if row["phrase"] in existing_phrases:
            raise SystemExit(f"Phrase already exists in source corpus: {row['phrase']}")

    translations = json.loads(EN_MAP.read_text(encoding="utf-8"))
    review = json.loads(REVIEW_JSON.read_text(encoding="utf-8"))
    overrides = json.loads(DISPATCH_OVERRIDES.read_text(encoding="utf-8"))
    selected_snapshot_fields = V4_FIELDS + ["source_id", "assigned_id"]
    selected_snapshot: list[dict[str, str]] = []

    for row in selected:
        source_theme, source_id, evidence, note = source_metadata(row)
        numeric_id = str(next_id)
        pid = f"sb_{numeric_id}"
        source_rows.append({
            "id": numeric_id,
            "phrase": row["phrase"],
            "approx_period": "2026",
            "theme": source_theme,
            "source_id": source_id,
            "evidence": evidence,
            "notes": note,
        })
        translations[pid] = TRANSLATIONS[row["phrase"]]
        review[pid] = editorial_record(row, pid, source_theme, source_id, evidence)
        override = dispatch_override_for(row, pid)
        if override:
            overrides[pid] = override
        selected_snapshot.append({**row, "source_id": source_id, "assigned_id": numeric_id})
        existing_phrases.add(row["phrase"])
        existing_ids.add(next_id)
        next_id += 1

    with SOURCE.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SOURCE_FIELDS)
        writer.writeheader()
        writer.writerows(source_rows)
    EN_MAP.write_text(json.dumps(translations, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    REVIEW_JSON.write_text(json.dumps(review, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    DISPATCH_OVERRIDES.write_text(json.dumps(overrides, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    with SELECTED_SNAPSHOT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=selected_snapshot_fields)
        writer.writeheader()
        writer.writerows(selected_snapshot)

    # Keep the human-readable review export aligned with the expanded source.
    review_fields = [
        "id", "phrase", "english", "sourceTheme", "evidence", "sourceId",
        "sourceCategory", "rightsStatus", "reviewBasis", "decision", "reviewCode",
        "lifecycle", "proposedPhrase", "proposedEnglish", "editorialTheme",
        "emotionTheme", "semanticCluster", "cadenceGroup", "englishStatus",
        "dispatchStatus", "dispatchConstraint", "reviewerNote", "reviewReason",
    ]
    with REVIEW_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=review_fields)
        writer.writeheader()
        for source_row in source_rows:
            pid = f"sb_{source_row['id']}"
            item = review[pid]
            writer.writerow({
                "id": pid,
                **{
                    field: json.dumps(item[field], ensure_ascii=False)
                    if field == "dispatchConstraint"
                    else item[field]
                    for field in review_fields
                    if field != "id"
                },
            })

    print(f"Promoted {len(selected)} selected candidates")
    print(f"Source rows: {len(source_rows)}")
    print(f"Assigned ids: sb_{next_id - len(selected)}..sb_{next_id - 1}")
    print(f"Hard dispatch overrides added: {sum(1 for row in selected if dispatch_override_for(row, ''))}")


if __name__ == "__main__":
    main()
